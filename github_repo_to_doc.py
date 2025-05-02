import os
import tempfile
import shutil
from git import Repo
from docx import Document
import platform
import subprocess

def clone_github_repo(repo_url):
    tmp_dir = tempfile.mkdtemp()
    Repo.clone_from(repo_url, tmp_dir)
    return tmp_dir

def extract_code_to_doc(repo_dir, output_file="output.docx", extensions=None):
    if extensions is None:
        extensions = ['.py', '.js', '.ts', '.html', '.css', '.java', '.c', '.cpp', '.sql']

    exclude_dirs = {'node_modules', 'venv', '.git', '__pycache__', 'build', 'dist', '.idea', '.vscode'}
    exclude_files = {'.DS_Store', 'package-lock.json', 'yarn.lock', 'requirements.txt'}  # optional

    doc = Document()
    doc.add_heading('Code Extracted from GitHub Repository', level=1)

    for root, dirs, files in os.walk(repo_dir):
        dirs[:] = [d for d in dirs if d not in exclude_dirs]

        for file in files:
            if file in exclude_files:
                continue
            if any(file.endswith(ext) for ext in extensions):
                full_path = os.path.join(root, file)
                try:
                    with open(full_path, 'r', encoding='utf-8', errors='ignore') as f:
                        code = f.read()
                        relative_path = os.path.relpath(full_path, repo_dir)
                        doc.add_heading(relative_path, level=2)
                        doc.add_paragraph(code)
                except Exception as e:
                    print(f"Failed to read {file}: {e}")

    doc.save(output_file)
    print(f"Document saved as {output_file}")


def open_docx(filepath):
    system = platform.system()
    if system == "Windows":
        os.startfile(filepath)
    elif system == "Darwin":  # macOS
        subprocess.run(["open", filepath])
    elif system == "Linux":
        subprocess.run(["xdg-open", filepath])
    else:
        print(f"Unsupported OS: {system}")

def main():
    repo_url = input("Enter GitHub repo URL: ").strip()
    repo_dir = clone_github_repo(repo_url)
    try:
        extract_code_to_doc(repo_dir)
        open_docx("output.docx")
    finally:
        shutil.rmtree(repo_dir)

if __name__ == "__main__":
    main()
